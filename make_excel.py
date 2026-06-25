import csv
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter

# ── Load CSV ────────────────────────────────────────────────────────────────
with open("imts_exhibitors.csv", encoding="utf-8-sig") as f:
    raw = list(csv.DictReader(f))

rows = []
for r in raw:
    if r.get("Company Name","").strip():
        rows.append({
            "Company Name": r.get("Company Name","").strip(),
            "Main Product":  r.get("Main Product","").strip(),
            "Country":       r.get("Country","").strip(),
            "Website":       r.get("Website","").strip(),
            "Email":         r.get("Email","").strip(),
        })

# ── Styles ───────────────────────────────────────────────────────────────────
DARK_BLUE   = "1F3864"
MID_BLUE    = "2E5999"
LIGHT_BLUE  = "D6E4F0"
ROW_ALT     = "EBF3FB"
WHITE       = "FFFFFF"
ACCENT      = "F0F4FB"

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, size=10, color="000000", name="Arial"):
    return Font(bold=bold, size=size, color=color, name=name)

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left(wrap=False):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

thin  = Side(style="thin",   color="B8CCE4")
thick = Side(style="medium", color="1F3864")

def border_thin():
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def border_thick():
    return Border(left=thick, right=thick, top=thick, bottom=thick)

# ── Workbook ─────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "IMTS 2026 Exhibitors"
ws.sheet_view.showGridLines = False

COLS = ["Company Name", "Main Product", "Country", "Website", "Email"]
COL_WIDTHS = [36, 52, 24, 38, 32]
DATA_START = 4   # row where data begins (1=title, 2=subtitle, 3=header)
TOTAL_ROWS = DATA_START + len(rows) - 1

# ── Row 1: Title ─────────────────────────────────────────────────────────────
ws.row_dimensions[1].height = 36
ws.merge_cells("A1:E1")
c = ws["A1"]
c.value = "IMTS 2026  |  Tooling & Workholding Exhibitors"
c.font      = Font(bold=True, size=16, color=WHITE, name="Arial")
c.fill      = fill(DARK_BLUE)
c.alignment = center()

# ── Row 2: Subtitle ───────────────────────────────────────────────────────────
ws.row_dimensions[2].height = 22
ws.merge_cells("A2:E2")
c = ws["A2"]
emails_found = sum(1 for r in rows if r["Email"])
c.value = (f"Total Exhibitors: {len(rows)}   |   Emails Found: {emails_found}   |   "
           f"Scraped: {date.today().strftime('%B %d, %Y')}")
c.font      = Font(size=9, color="1F3864", name="Arial")
c.fill      = fill(LIGHT_BLUE)
c.alignment = center()

# ── Row 3: Header ─────────────────────────────────────────────────────────────
ws.row_dimensions[3].height = 24
for col_idx, col_name in enumerate(COLS, 1):
    c = ws.cell(row=3, column=col_idx, value=col_name.upper())
    c.font      = Font(bold=True, size=10, color=WHITE, name="Arial")
    c.fill      = fill(MID_BLUE)
    c.alignment = center()
    c.border    = border_thin()

# ── Data Rows ─────────────────────────────────────────────────────────────────
for row_num, row_data in enumerate(rows, DATA_START):
    ws.row_dimensions[row_num].height = 40
    bg = WHITE if (row_num - DATA_START) % 2 == 0 else ROW_ALT

    for col_idx, col_name in enumerate(COLS, 1):
        val = row_data[col_name]
        c   = ws.cell(row=row_num, column=col_idx)
        c.fill   = fill(bg)
        c.border = border_thin()

        if col_name == "Company Name":
            c.value     = val
            c.font      = Font(bold=True, size=10, color="1F3864", name="Arial")
            c.alignment = left()

        elif col_name == "Main Product":
            c.value     = val
            c.font      = Font(size=9, color="333333", name="Arial")
            c.alignment = left(wrap=True)

        elif col_name == "Country":
            c.value     = val
            c.font      = font(size=10)
            c.alignment = center()

        elif col_name == "Website":
            if val:
                c.value     = val
                c.hyperlink = val
                c.font      = Font(size=9, color="1155CC", underline="single",
                                   name="Arial")
            else:
                c.value = ""
                c.font  = font(size=9)
            c.alignment = left()

        elif col_name == "Email":
            if val:
                c.value     = val
                c.hyperlink = f"mailto:{val}"
                c.font      = Font(size=9, color="1155CC", underline="single",
                                   name="Arial")
            else:
                c.value = ""
                c.font  = font(size=9)
            c.alignment = left()

# ── Column Widths ─────────────────────────────────────────────────────────────
for col_idx, width in enumerate(COL_WIDTHS, 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# ── Thick outer border around whole table ────────────────────────────────────
for row in ws.iter_rows(min_row=1, max_row=TOTAL_ROWS, min_col=1, max_col=5):
    for cell in row:
        old = cell.border
        new_left   = thick if cell.column == 1 else old.left
        new_right  = thick if cell.column == 5 else old.right
        new_top    = thick if cell.row    == 1 else old.top
        new_bottom = thick if cell.row    == TOTAL_ROWS else old.bottom
        cell.border = Border(left=new_left, right=new_right,
                             top=new_top,   bottom=new_bottom)

# ── Freeze panes & auto-filter ────────────────────────────────────────────────
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:E3"

# ── Save ──────────────────────────────────────────────────────────────────────
out = "imts_exhibitors_formatted.xlsx"
wb.save(out)
print(f"Saved: {out}  ({len(rows)} rows, {emails_found} emails)")
